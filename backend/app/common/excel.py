"""Excel 导入导出工具（openpyxl）。"""

import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from typing import Any
from urllib.parse import quote

from fastapi import UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.errors import BizError


def _cell(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    if isinstance(v, (list, dict)):
        return str(v)
    return v


def build_workbook(sheet_title: str, columns: Sequence[tuple[str, str]], rows: Iterable[dict]) -> bytes:
    """columns: [(字段, 表头)]"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for idx, (_, title) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=idx, value=title)
        c.font = Font(bold=True)
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(idx)].width = max(12, min(40, len(title) * 2 + 4))
    for r, row in enumerate(rows, start=2):
        for idx, (key, _) in enumerate(columns, start=1):
            ws.cell(row=r, column=idx, value=_cell(row.get(key)))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(content: bytes, filename: str) -> StreamingResponse:
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def export_xlsx(filename: str, columns: Sequence[tuple[str, str]], rows: Iterable[dict]) -> StreamingResponse:
    return xlsx_response(build_workbook(filename.rsplit(".", 1)[0], columns, rows), filename)


def read_upload(file: UploadFile, columns: Sequence[tuple[str, str]], *, max_rows: int = 20000) -> list[dict]:
    """读取上传的 xlsx，按表头（中文表头或字段名均可）映射为 dict 列表。"""
    data = file.file.read()
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise BizError("无法解析 Excel 文件，请使用 .xlsx 格式") from exc
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration as exc:
        raise BizError("Excel 文件为空") from exc
    title_to_key = {}
    for key, title in columns:
        title_to_key[title.strip()] = key
        title_to_key[key] = key
    keys = [title_to_key.get(str(h).strip()) if h is not None else None for h in header]
    if not any(keys):
        raise BizError("未识别到有效表头，请下载导入模板")
    result = []
    for i, row in enumerate(rows_iter, start=2):
        if i > max_rows + 1:
            raise BizError(f"单次最多导入 {max_rows} 行")
        if row is None or all(v in (None, "") for v in row):
            continue
        item = {"_row": i}
        for k, v in zip(keys, row, strict=False):
            if k:
                if isinstance(v, str):
                    v = v.strip()
                if isinstance(v, datetime):
                    v = v.date() if v.time() == datetime.min.time() else v
                item[k] = v
        result.append(item)
    return result


def template_response(filename: str, columns: Sequence[tuple[str, str]], sample: dict | None = None) -> StreamingResponse:
    return export_xlsx(filename, columns, [sample] if sample else [])


def as_date(v: Any) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v)[:10])
